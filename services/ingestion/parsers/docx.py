"""DOCX 解析器（``.docx``，基于 python-docx）。

保留标题层级与表格
------------------
这是计划 2.2 对本格式的两条明确要求，实现要点：

* **标题层级**：Word 的层级信息有三个来源，按可靠性依次取
  ``paragraph.style.name``（``Heading 1`` / ``标题 1``）→ ``style.style_id``（``Heading1``）
  → 段落 XML 里的 ``w:outlineLvl``。多路兜底是因为「标题 1」是本地化样式名，
  而某些模板（尤其国内 WPS 导出的）只写 outlineLvl，不写样式名。
* **表格**：``w:tbl`` 按文档顺序穿插在段落之间，因此不能只遍历
  ``document.paragraphs``（它**看不到表格**）。这里直接按 ``w:body`` 的子元素顺序走，
  段落与表格混排时顺序才与原文一致。

表格首行按表头处理，零宽行会被 :func:`build_table_json` 补齐。
单元格合并（``w:gridSpan`` / ``w:vMerge``）不做特殊还原，合并区域的文字会在
参与合并的每个格子里各出现一次 —— 这是有意的取舍：还原合并需要处理
vMerge 的 continue/restart 状态机，而错误还原会把内容**错位到别的列**，
比重复出现更危险（重复只是检索时多命中一次）。

图片以 ``docx:<rId>`` 形式记录引用（走 OOXML 的关系 ID），不抽取二进制。
"""

from __future__ import annotations

import io
import re
from typing import Any, Iterator

from services.ingestion.parsers._text import first_nonempty_line
from services.ingestion.parsers.base import (
    BLOCK_HEADING,
    BLOCK_IMAGE,
    BLOCK_LIST,
    BLOCK_PARAGRAPH,
    BLOCK_TABLE,
    ERROR_CORRUPT_DOCUMENT,
    ERROR_MISSING_DEPENDENCY,
    Block,
    ParsedDocument,
    ParserError,
    TemplateDocumentParser,
    build_table_json,
    render_table_text,
)
from services.ingestion.parsers.ocr import OcrEngine, load_default_engine


_HEADING_NAME_RE = re.compile(r"(?:Heading|标题|Titre|Überschrift)\s*(\d+)", re.IGNORECASE)
_HEADING_ID_RE = re.compile(r"Heading(\d+)", re.IGNORECASE)
_LIST_NAME_HINTS: tuple[str, ...] = ("List", "列表", "Bullet")
_TITLE_STYLE_NAMES: frozenset[str] = frozenset({"Title", "标题"})


def _require_python_docx() -> tuple[Any, Any, Any, Any]:
    """延迟导入 python-docx，缺失时给结构化错误。"""

    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise ParserError(
            ERROR_MISSING_DEPENDENCY,
            "未安装 python-docx，无法解析 DOCX",
            parser_name="docx",
            detail={"missing_package": "python-docx"},
        ) from exc
    return Document, Table, Paragraph, qn


def _iter_body_items(body: Any, parent: Any, qn: Any, paragraph_type: Any, table_type: Any) -> Iterator[Any]:
    """按文档顺序产出段落与表格，并下钻到内容控件（``w:sdt``）内部。

    只遍历 ``document.paragraphs`` 会漏掉全部表格；只遍历 ``w:body`` 的直接子元素
    会漏掉被内容控件包起来的段落。两种都属于「静默丢内容」，必须一起处理。
    """

    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield paragraph_type(child, parent)
        elif child.tag == qn("w:tbl"):
            yield table_type(child, parent)
        elif child.tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from _iter_body_items(content, parent, qn, paragraph_type, table_type)


def _heading_level(paragraph: Any, qn: Any) -> int | None:
    style = getattr(paragraph, "style", None)
    name = (getattr(style, "name", "") or "").strip() if style is not None else ""
    if name in _TITLE_STYLE_NAMES:
        return 1

    match = _HEADING_NAME_RE.search(name)
    if match:
        return int(match.group(1))

    style_id = (getattr(style, "style_id", "") or "") if style is not None else ""
    match = _HEADING_ID_RE.search(style_id)
    if match:
        return int(match.group(1))

    properties = paragraph._p.pPr
    if properties is not None:
        outline = properties.find(qn("w:outlineLvl"))
        if outline is not None:
            value = outline.get(qn("w:val"))
            if value is not None:
                try:
                    return int(value) + 1
                except ValueError:
                    return None
    return None


def _is_list_item(paragraph: Any, qn: Any) -> bool:
    properties = paragraph._p.pPr
    if properties is not None and properties.find(qn("w:numPr")) is not None:
        return True
    style = getattr(paragraph, "style", None)
    name = (getattr(style, "name", "") or "") if style is not None else ""
    return any(hint in name for hint in _LIST_NAME_HINTS)


def _image_refs(paragraph: Any, qn: Any, index: int) -> list[str]:
    """取出段落下所有图片的关系 ID，作为图片引用。"""

    references: list[str] = []
    for drawing in paragraph._p.findall(f".//{qn('w:drawing')}"):
        for blip in drawing.findall(f".//{qn('a:blip')}"):
            embed = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
            if embed:
                references.append(f"docx:{embed}")
    for _ in paragraph._p.findall(f".//{qn('w:pict')}"):
        references.append(f"docx:pict{index}")
    return references


class DocxParser(TemplateDocumentParser):
    supported_types = ("docx",)
    parser_name = "docx"
    parser_version = "1.0.0"

    def __init__(self, ocr_engine: OcrEngine | None = None, use_default_ocr: bool = True) -> None:
        # 与 PDF 对齐：只有「全文没有文本」的 DOCX（整页都是图片）才值得 OCR。
        # 这里不读阈值环境变量，因为 DOCX 的文本层要么在要么不在，几乎不存在中间态。
        if ocr_engine is not None:
            self.ocr_engine: OcrEngine | None = ocr_engine
        elif use_default_ocr:
            self.ocr_engine = load_default_engine()
        else:
            self.ocr_engine = None

    def _parse(self, source: bytes, filename: str) -> ParsedDocument:
        document_type, table_type, paragraph_type, qn = _require_python_docx()

        try:
            document = document_type(io.BytesIO(source))
        except Exception as exc:
            raise ParserError(
                ERROR_CORRUPT_DOCUMENT,
                f"DOCX 无法打开：{type(exc).__name__}: {exc}",
                parser_name=self.parser_name,
                filename=filename,
                detail={"exception": type(exc).__name__},
            ) from exc

        blocks: list[Block] = []
        heading_stack: list[tuple[int, str]] = []
        stats = {"heading": 0, "list": 0, "table": 0, "image": 0, "paragraph": 0}

        for index, item in enumerate(_iter_body_items(document.element.body, document, qn, paragraph_type, table_type)):
            path = tuple(text for _, text in heading_stack)

            if isinstance(item, table_type):
                table_json = self._table_json(item)
                if table_json is not None:
                    blocks.append(
                        Block(
                            type=BLOCK_TABLE,
                            text=render_table_text(table_json),
                            heading_path=path,
                            table_json=table_json,
                        )
                    )
                    stats["table"] += 1
                continue

            text = (item.text or "").strip()
            for reference in _image_refs(item, qn, index):
                blocks.append(Block(type=BLOCK_IMAGE, text="", heading_path=path, image_ref=reference))
                stats["image"] += 1

            if not text:
                continue

            level = _heading_level(item, qn)
            if level is not None and level <= 6:
                while heading_stack and heading_stack[-1][0] >= level:
                    heading_stack.pop()
                heading_stack.append((level, text))
                blocks.append(
                    Block(type=BLOCK_HEADING, text=text, heading_path=tuple(t for _, t in heading_stack))
                )
                stats["heading"] += 1
                continue

            if _is_list_item(item, qn):
                blocks.append(Block(type=BLOCK_LIST, text=text, heading_path=path))
                stats["list"] += 1
            else:
                blocks.append(Block(type=BLOCK_PARAGRAPH, text=text, heading_path=path))
                stats["paragraph"] += 1

        if not any(block.text.strip() for block in blocks):
            raise ParserError(
                ERROR_CORRUPT_DOCUMENT,
                "DOCX 没有可提取的文本内容（可能整份文档都是图片）",
                parser_name=self.parser_name,
                filename=filename,
                detail={"block_count": len(blocks)},
            )

        title = ""
        core_title = (document.core_properties.title or "").strip()
        if core_title:
            title = core_title
        if not title:
            title = next((block.text for block in blocks if block.type == BLOCK_HEADING), "")
        if not title:
            title = first_nonempty_line(next((block.text for block in blocks if block.text), ""))

        metadata: dict[str, Any] = {
            "docx_paragraph_count": stats["paragraph"],
            "docx_heading_count": stats["heading"],
            "docx_list_count": stats["list"],
            "docx_table_count": stats["table"],
            "docx_image_count": stats["image"],
            "docx_character_count": sum(len(block.text) for block in blocks),
        }
        return ParsedDocument(title=title, blocks=tuple(blocks), metadata=metadata)

    def _table_json(self, table: Any) -> dict[str, Any] | None:
        """把 DOCX 表格转成稳定的表格 JSON（首行作表头）。"""

        rows = [[(cell.text or "").strip() for cell in row.cells] for row in table.rows]
        rows = [row for row in rows if any(cell for cell in row)]
        if not rows:
            return None
        return build_table_json(rows[0], rows[1:])


__all__ = ["DocxParser"]
